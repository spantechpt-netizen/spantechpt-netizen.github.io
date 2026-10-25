"""BBS workbook with formulas from the generators' exports (out/*.bbs.json, written by BarList.export).

    python3 bbs_xlsx.py BBS.xlsx out/FOOTINGS.bbs.json out/NECKS.bbs.json [...]

One sheet per ELEMENT (each footing type, each combined footing / raft, each column-neck model - engineer,
25-10-2026), titled with its real name and the number of elements it covers. One row per bar mark: drawing number,
mark, diameter, shape + segment lengths, quantity per element, number of elements carrying the mark, total quantity
(= per element x elements, or the counted total when the quantity varies per element: neck ties), cutting length,
total length, unit weight d^2/162 and weight - all as Excel formulas, so editing a quantity or a length updates the
sheet and the summary. 'Summary' totals every diameter per element and per group with SUMIFS + the grand total.
Net engineering quantities: no waste allowance.
"""
import json, math, re, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SHAPE = {'S': 'STRAIGHT', 'U': 'U-BAR', 'L': 'L-BAR', 'POLY': 'TIE (2 HOOKS 135)', 'HT': 'STIRRUP (2 HOOKS 135)',
         'ST': 'CLOSED LOOP (LAPPED)', 'CH': 'CHAIR', 'OPEN': 'OPEN BAR', 'PROF': 'BENT BAR (PROFILE)'}


def group_of(el):
    if ' ON ' in el or re.match(r'^C[\d\- ]', el.upper()) and not el.upper().startswith('CF'): return 'Column necks'
    if el.upper().startswith('CF'): return 'Combined footings'
    if el.upper().startswith('RAFT'): return 'Rafts'
    if el.upper().startswith('F'): return 'Isolated footings'
    return 'Other'


def segs(shape):
    kind, rest = shape[0], shape[1:]
    if kind == 'POLY':
        pts = [tuple(p) for p in rest]
        return [round(math.dist(pts[i], pts[(i + 1) % len(pts)])) for i in range(len(pts))][:6]
    return [v for v in rest if isinstance(v, (int, float))][:6]


def sheet_name(el, used):
    n = re.sub(r'[\[\]\*\?/\\:]', 'x', el)[:28]
    k, b = 2, n
    while n in used: n = f'{b[:25]}~{k}'; k += 1
    used.add(n); return n


def main(out, files):
    elems = {}                                     # element -> dict(dwg, n, rows)
    for fn in files:
        d = json.load(open(fn))
        for r in d['rows']:
            for el, q in r.get('use', {}).items():
                e = elems.setdefault(el, dict(dwg=d['dwg'].get(el, ''), n=d.get('counts', {}).get(el, 1), rows=[], src=fn))
                if r.get('split', {}).get(el):                # exact counts per group (neck ties per footing depth)
                    for n_c, ncol, lab in r['split'][el]:
                        e['rows'].append(dict(mark=r['mark'], d=r['d'], shape=r['shape'], L=r['L'], q=n_c * ncol, ne=ncol, layer=lab))
                    continue
                ne = r.get('use_n', {}).get(el) or e['n']
                e['rows'].append(dict(mark=r['mark'], d=r['d'], shape=r['shape'], L=r['L'], q=q, ne=ne, layer=r.get('layer', '')))
    hdr = Font(bold=True, color='FFFFFF'); fill = PatternFill('solid', fgColor='1D3557'); thin = Side(style='thin', color='999999')
    bd = Border(top=thin, bottom=thin, left=thin, right=thin)
    wb = Workbook(); summ = wb.active; summ.title = 'Summary'
    head = ['Bar mark', 'Dia (mm)', 'Shape', 'A', 'B', 'C', 'D', 'E', 'F', 'Qty per element', 'No. of elements', 'Total qty',
            'Cutting length (mm)', 'Total length (m)', 'Unit weight (kg/m) = d²/162', 'Total weight (kg)', 'Layer / note']
    used, sheets = {'Summary'}, []
    order = sorted(elems, key=lambda e: (['Isolated footings', 'Combined footings', 'Rafts', 'Column necks', 'Other'].index(group_of(e)), e))
    for el in order:
        e = elems[el]; nm = sheet_name(el, used); ws = wb.create_sheet(nm); sheets.append((el, nm, e))
        ws['A1'] = el; ws['A1'].font = Font(bold=True, size=14)
        ws['A2'] = f"Drawing: {e['dwg']}"; ws['F2'] = 'No. of elements covered:'; ws['I2'] = e['n']; ws['I2'].font = Font(bold=True)
        ws['K2'] = 'Group:'; ws['L2'] = group_of(el)
        ws.append([]); ws.append(head)
        for c in ws[4]: c.font = hdr; c.fill = fill; c.alignment = Alignment(wrap_text=True, horizontal='center', vertical='center')
        for i, r in enumerate(sorted(e['rows'], key=lambda r: r['mark']), 5):
            sg = segs(r['shape']) + [None] * 6
            per = r['q'] / r['ne']
            exact = abs(per - round(per)) < 1e-6
            ws.append([f"{r['mark']:02d}", r['d'], SHAPE.get(r['shape'][0], r['shape'][0])] + sg[:6] +
                      [round(per) if exact else 'VAR', r['ne'], f'=J{i}*K{i}' if exact else r['q'], r['L'], f'=L{i}*M{i}/1000',
                       f'=B{i}^2/162', f'=N{i}*O{i}', r['layer'] + ('' if exact else ' (qty varies per element: counted total)')])
            for col_, fm in (('N', '0.00'), ('O', '0.000'), ('P', '#,##0.00')): ws[f'{col_}{i}'].number_format = fm
        n = ws.max_row
        ws.append(['TOTAL'] + [''] * 12 + [f'=SUM(N5:N{n})', '', f'=SUM(P5:P{n})'])
        for c in ws[n + 1]: c.font = Font(bold=True)
        ws[f'P{n + 1}'].number_format = '#,##0.00'; ws[f'N{n + 1}'].number_format = '0.00'
        for row in ws.iter_rows(min_row=4):
            for c in row: c.border = bd
        for j, w in enumerate([8, 8, 22, 7, 7, 7, 7, 7, 7, 10, 10, 10, 12, 12, 14, 14, 30], 1):
            ws.column_dimensions[get_column_letter(j)].width = w
        ws.freeze_panes = 'A5'
    # summary: one row per element, one column per diameter (SUMIFS on each sheet) + group / grand totals
    dias = sorted({r['d'] for e in elems.values() for r in e['rows']})
    summ.append(['STEEL SUMMARY BY DIAMETER (kg) - net quantities, no waste allowance'])
    summ['A1'].font = Font(bold=True, size=13)
    summ.append([])
    summ.append(['Element', 'Group', 'No. of elements', 'Drawing'] + [f'T{d}' for d in dias] + ['Total (kg)'])
    for c in summ[3]: c.font = hdr; c.fill = fill; c.alignment = Alignment(wrap_text=True, horizontal='center')
    first = 4
    for k, (el, nm, e) in enumerate(sheets):
        i = first + k
        cells = [f"=SUMIFS('{nm}'!P:P,'{nm}'!B:B,{d})" for d in dias]
        last = get_column_letter(4 + len(dias))
        summ.append([el, group_of(el), e['n'], e['dwg']] + cells + [f'=SUM(E{i}:{last}{i})'])
    lastrow = summ.max_row
    for g in ['Isolated footings', 'Combined footings', 'Rafts', 'Column necks', 'Other']:
        if not any(group_of(el) == g for el, _, _ in sheets): continue
        i = summ.max_row + 1
        summ.append([f'TOTAL {g.upper()}', g, '', ''] + [f'=SUMIFS({get_column_letter(5 + j)}{first}:{get_column_letter(5 + j)}{lastrow},$B${first}:$B${lastrow},"{g}")' for j in range(len(dias) + 1)])
        for c in summ[i]: c.font = Font(bold=True)
    i = summ.max_row + 1
    summ.append(['GRAND TOTAL (kg)', '', '', ''] + [f'=SUM({get_column_letter(5 + j)}{first}:{get_column_letter(5 + j)}{lastrow})' for j in range(len(dias) + 1)])
    for c in summ[i]: c.font = Font(bold=True)
    summ.append(['GRAND TOTAL (t)', '', '', '', ''] + [''] * (len(dias) - 1) + [f'={get_column_letter(5 + len(dias))}{i}/1000'])
    for row in summ.iter_rows(min_row=3):
        for c in row:
            c.border = bd
            if c.column > 4: c.number_format = '#,##0.00'
    for j, w in enumerate([30, 18, 10, 34] + [11] * (len(dias) + 1), 1): summ.column_dimensions[get_column_letter(j)].width = w
    summ.cell(summ.max_row + 2, 1, 'Unit weight = d²/162 (kg/m); weight = quantity x cutting length (m) x unit weight. Laps are inside the cutting lengths (never added twice).')
    wb.save(out)
    print(out, len(sheets), 'element sheets', dias)


if __name__ == '__main__':
    main(sys.argv[1], sys.argv[2:])
